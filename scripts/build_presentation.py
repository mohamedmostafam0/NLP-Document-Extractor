"""Generate the Docxtract project presentation (Docxtract.pptx).

Run from the repo root:
    python scripts/build_presentation.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Theme — dark mode that matches the Docxtract UI
# ---------------------------------------------------------------------------
BG       = RGBColor(0x0F, 0x17, 0x23)
PANEL    = RGBColor(0x18, 0x22, 0x32)
ACCENT   = RGBColor(0x4D, 0x9D, 0xFF)
ACCENT_2 = RGBColor(0x2D, 0xD4, 0xBF)
WARN     = RGBColor(0xF5, 0xA6, 0x24)
TEXT     = RGBColor(0xE6, 0xED, 0xF5)
TEXT_DIM = RGBColor(0x9A, 0xA8, 0xBD)
DIVIDER  = RGBColor(0x2A, 0x36, 0x4C)

REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT    = REPO_ROOT / "Docxtract.pptx"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _set_bg(slide, color=BG):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                Inches(13.333), Inches(7.5))
    bg.fill.solid(); bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    return bg


def _add_text(slide, left, top, width, height, text, *,
              font_size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT,
              font_name="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return box


def _add_bullets(slide, left, top, width, height, items, *,
                 font_size=20, color=TEXT, bullet_color=ACCENT,
                 line_spacing=1.5):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        p.space_after = Pt(8)
        b = p.add_run()
        b.text = "▸  "
        b.font.size = Pt(font_size)
        b.font.color.rgb = bullet_color
        b.font.bold = True
        b.font.name = "Calibri"
        r = p.add_run()
        r.text = item
        r.font.size = Pt(font_size)
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return box


def _add_chip(slide, left, top, text, color=ACCENT, width=None):
    if width is None:
        width = Inches(max(1.2, 0.16 * len(text) + 0.35))
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top,
                                  width, Inches(0.4))
    chip.fill.solid(); chip.fill.fore_color.rgb = PANEL
    chip.line.color.rgb = color
    chip.line.width = Pt(1.25)
    tf = chip.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return chip


def _add_card(slide, left, top, width, height, *, border=ACCENT):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  left, top, width, height)
    card.fill.solid(); card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = border
    card.line.width = Pt(1.0)
    card.shadow.inherit = False
    return card


def _add_divider(slide, left, top, width):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                  left, top, width, Inches(0.04))
    line.fill.solid(); line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()
    return line


def _slide_header(slide, kicker: str, title: str):
    _add_text(slide, Inches(0.7), Inches(0.45), Inches(8), Inches(0.4),
              kicker.upper(), font_size=13, bold=True, color=ACCENT)
    _add_text(slide, Inches(0.7), Inches(0.8), Inches(12), Inches(0.9),
              title, font_size=40, bold=True, color=TEXT)
    _add_divider(slide, Inches(0.7), Inches(1.75), Inches(1.4))


def _footer(slide, page_num: int, total: int):
    _add_text(slide, Inches(0.7), Inches(7.05), Inches(8), Inches(0.3),
              "Docxtract", font_size=10, color=TEXT_DIM)
    _add_text(slide, Inches(11.5), Inches(7.05), Inches(1.3), Inches(0.3),
              f"{page_num:02d} / {total:02d}",
              font_size=10, color=TEXT_DIM, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def make_title_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                 Inches(0.3), Inches(7.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    _add_text(slide, Inches(0.9), Inches(1.2), Inches(11), Inches(0.5),
              "SELECTED TOPICS IN DATA SCIENCE  ·  2026",
              font_size=14, bold=True, color=ACCENT)

    _add_text(slide, Inches(0.9), Inches(2.2), Inches(12), Inches(1.5),
              "Docxtract",
              font_size=96, bold=True, color=TEXT)

    _add_text(slide, Inches(0.9), Inches(3.7), Inches(12), Inches(0.8),
              "NLP Document Intelligence Pipeline",
              font_size=30, color=ACCENT_2)

    _add_divider(slide, Inches(0.9), Inches(4.7), Inches(1.4))

    _add_text(slide, Inches(0.9), Inches(4.95), Inches(11), Inches(0.5),
              "Unstructured documents → validated JSON.",
              font_size=20, color=TEXT_DIM)


def make_problem_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "01  ·  Motivation", "The Problem")

    _add_text(slide, Inches(0.7), Inches(2.1), Inches(12), Inches(0.6),
              "Business data is trapped inside documents.",
              font_size=24, color=TEXT_DIM)

    cards = [
        ("Slow",    "Manual data entry doesn't scale."),
        ("Messy",   "Every layout is different."),
        ("Noisy",   "OCR alone gives you raw text — not data."),
    ]
    x = Inches(0.7)
    for title, body in cards:
        _add_card(slide, x, Inches(3.3), Inches(4.0), Inches(2.5), border=ACCENT)
        _add_text(slide, x + Inches(0.3), Inches(3.6), Inches(3.4), Inches(0.6),
                  title, font_size=26, bold=True, color=ACCENT)
        _add_text(slide, x + Inches(0.3), Inches(4.3), Inches(3.4), Inches(1.5),
                  body, font_size=15, color=TEXT)
        x += Inches(4.2)

    _add_text(slide, Inches(0.7), Inches(6.2), Inches(12), Inches(0.5),
              "Docxtract automates the whole flow.",
              font_size=18, bold=True, color=ACCENT_2)

    _footer(slide, page, total)


def make_features_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "02  ·  Overview", "Key Features")

    _add_bullets(slide, Inches(0.7), Inches(2.2), Inches(12), Inches(5), [
        "Multi-model NLP pipeline",
        "Real-time progress tracking",
        "Human-in-the-loop review",
        "JSON & CSV export",
        "Premium dark-mode UI",
        "Fully containerized",
    ], font_size=24, line_spacing=1.6)

    _footer(slide, page, total)


def make_architecture_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "03  ·  System Design", "Architecture")

    # Frontend
    _add_card(slide, Inches(0.7), Inches(2.2), Inches(12), Inches(0.8),
              border=ACCENT_2)
    _add_text(slide, Inches(0.7), Inches(2.35), Inches(12), Inches(0.55),
              "Frontend SPA",
              font_size=20, bold=True, color=ACCENT_2, align=PP_ALIGN.CENTER)

    # Backend
    _add_card(slide, Inches(0.7), Inches(3.2), Inches(12), Inches(0.8),
              border=ACCENT)
    _add_text(slide, Inches(0.7), Inches(3.35), Inches(12), Inches(0.55),
              "FastAPI Backend",
              font_size=20, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

    # Engine
    boxes = [("OCR", "Tesseract"), ("NER", "spaCy + Transformer"), ("Regex", "Patterns")]
    x = Inches(0.7)
    for title, sub in boxes:
        _add_card(slide, x, Inches(4.2), Inches(3.95), Inches(1.0),
                  border=ACCENT)
        _add_text(slide, x, Inches(4.35), Inches(3.95), Inches(0.4),
                  title, font_size=18, bold=True, color=ACCENT,
                  align=PP_ALIGN.CENTER)
        _add_text(slide, x, Inches(4.75), Inches(3.95), Inches(0.4),
                  sub, font_size=12, color=TEXT_DIM, align=PP_ALIGN.CENTER)
        x += Inches(4.05)

    # Storage
    _add_card(slide, Inches(0.7), Inches(5.4), Inches(5.9), Inches(1.0),
              border=WARN)
    _add_text(slide, Inches(0.7), Inches(5.65), Inches(5.9), Inches(0.55),
              "MinIO — files",
              font_size=18, bold=True, color=WARN, align=PP_ALIGN.CENTER)

    _add_card(slide, Inches(6.8), Inches(5.4), Inches(5.9), Inches(1.0),
              border=WARN)
    _add_text(slide, Inches(6.8), Inches(5.65), Inches(5.9), Inches(0.55),
              "PostgreSQL — data",
              font_size=18, bold=True, color=WARN, align=PP_ALIGN.CENTER)

    _footer(slide, page, total)


def make_pipeline_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "04  ·  The Core", "6-Phase Pipeline")

    phases = [
        ("01", "Upload",        ACCENT),
        ("02", "Ingestion",     ACCENT),
        ("03", "Preprocessing", ACCENT_2),
        ("04", "Extraction",    ACCENT_2),
        ("05", "Mapping",       WARN),
        ("06", "Validation",    WARN),
    ]
    x0, y0 = Inches(0.7), Inches(2.6)
    w, h   = Inches(2.0), Inches(2.4)
    gap    = Inches(0.05)

    for i, (num, title, color) in enumerate(phases):
        x = x0 + i * (w + gap)
        _add_card(slide, x, y0, w, h, border=color)
        _add_text(slide, x, y0 + Inches(0.4), w, Inches(0.7),
                  num, font_size=44, bold=True, color=color,
                  align=PP_ALIGN.CENTER)
        _add_text(slide, x, y0 + Inches(1.4), w, Inches(0.5),
                  title, font_size=15, bold=True, color=TEXT,
                  align=PP_ALIGN.CENTER)

    _add_text(slide, Inches(0.7), Inches(5.5), Inches(12), Inches(0.5),
              "Live progress via Server-Sent Events.",
              font_size=18, color=ACCENT_2, align=PP_ALIGN.CENTER)

    _footer(slide, page, total)


def make_phase_details_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "05  ·  Deep Dive", "What Each Phase Does")

    rows = [
        ("Ingestion",     "PDF · DOCX · OCR"),
        ("Preprocessing", "Clean · dedupe · detect language"),
        ("Extraction",    "spaCy + Transformer + regex"),
        ("Mapping",       "Project onto typed schema"),
        ("Validation",    "Required fields + QC checks"),
    ]

    y = Inches(2.1)
    for name, sub in rows:
        _add_text(slide, Inches(0.8), y, Inches(4), Inches(0.6),
                  name, font_size=22, bold=True, color=ACCENT)
        _add_text(slide, Inches(5.0), y + Inches(0.07), Inches(8),
                  Inches(0.6), sub, font_size=18, color=TEXT)
        y += Inches(0.85)

    _footer(slide, page, total)


def make_schemas_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "06  ·  Domain Models", "Supported Documents")

    docs = [
        ("Business Card",  ACCENT,   "name · email · phone · company"),
        ("Resume",         ACCENT_2, "name · email · skills · experience"),
        ("Medical Report", WARN,     "patient · visit date · diagnoses"),
    ]

    y = Inches(2.4)
    for title, color, fields in docs:
        _add_card(slide, Inches(0.7), y, Inches(12), Inches(1.2),
                  border=color)
        _add_text(slide, Inches(1.0), y + Inches(0.2), Inches(4.5),
                  Inches(0.6), title, font_size=22, bold=True, color=color)
        _add_text(slide, Inches(5.5), y + Inches(0.32), Inches(7),
                  Inches(0.6), fields, font_size=15, color=TEXT_DIM)
        y += Inches(1.4)

    _footer(slide, page, total)


def make_lifecycle_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "07  ·  Lifecycle", "Two Outcomes")

    _add_card(slide, Inches(0.9), Inches(2.7), Inches(5.7), Inches(3.0),
              border=ACCENT_2)
    _add_text(slide, Inches(0.9), Inches(3.0), Inches(5.7), Inches(0.8),
              "✓  Completed", font_size=32, bold=True, color=ACCENT_2,
              align=PP_ALIGN.CENTER)
    _add_text(slide, Inches(1.1), Inches(4.0), Inches(5.3), Inches(1.8),
              "All required fields present.\nHigh confidence.",
              font_size=18, color=TEXT, align=PP_ALIGN.CENTER)

    _add_card(slide, Inches(6.8), Inches(2.7), Inches(5.7), Inches(3.0),
              border=WARN)
    _add_text(slide, Inches(6.8), Inches(3.0), Inches(5.7), Inches(0.8),
              "⚠  Needs Review", font_size=32, bold=True, color=WARN,
              align=PP_ALIGN.CENTER)
    _add_text(slide, Inches(7.0), Inches(4.0), Inches(5.3), Inches(1.8),
              "Missing fields or low confidence.\nHuman corrects in UI.",
              font_size=18, color=TEXT, align=PP_ALIGN.CENTER)

    _footer(slide, page, total)


def make_results_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "08  ·  Results", "Input → Output")

    _add_text(slide, Inches(0.7), Inches(2.0), Inches(6), Inches(0.5),
              "Input", font_size=18, bold=True, color=ACCENT)
    _add_card(slide, Inches(0.7), Inches(2.5), Inches(6), Inches(4.0),
              border=DIVIDER)
    raw = (
        "Richard Hendricks\n"
        "CEO & Founder\n"
        "Pied Piper Inc.\n"
        "(415) 555-0142\n"
        "richard@piedpiper.com"
    )
    _add_text(slide, Inches(0.95), Inches(2.75), Inches(5.5), Inches(3.5),
              raw, font_size=15, color=TEXT, font_name="Consolas")

    _add_text(slide, Inches(7.0), Inches(2.0), Inches(6), Inches(0.5),
              "Output", font_size=18, bold=True, color=ACCENT_2)
    _add_card(slide, Inches(7.0), Inches(2.5), Inches(5.9), Inches(4.0),
              border=ACCENT_2)
    out = (
        '{\n'
        '  "name":    "Richard Hendricks",\n'
        '  "title":   "CEO & Founder",\n'
        '  "company": "Pied Piper Inc.",\n'
        '  "phone":   "+14155550142",\n'
        '  "email":   "richard@piedpiper.com"\n'
        '}'
    )
    _add_text(slide, Inches(7.2), Inches(2.75), Inches(5.6), Inches(3.5),
              out, font_size=14, color=TEXT, font_name="Consolas")

    _footer(slide, page, total)


def make_testing_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "09  ·  Quality", "Testing")

    stats = [
        ("81",   "tests",          ACCENT),
        ("9",    "modules",        ACCENT_2),
        ("100%", "pass rate",      WARN),
    ]
    x = Inches(1.2)
    for num, label, color in stats:
        _add_card(slide, x, Inches(2.4), Inches(3.4), Inches(2.6),
                  border=color)
        _add_text(slide, x, Inches(2.7), Inches(3.4), Inches(1.2),
                  num, font_size=72, bold=True, color=color,
                  align=PP_ALIGN.CENTER)
        _add_text(slide, x, Inches(4.2), Inches(3.4), Inches(0.6),
                  label, font_size=18, color=TEXT,
                  align=PP_ALIGN.CENTER)
        x += Inches(3.6)

    _add_text(slide, Inches(0.7), Inches(5.6), Inches(12), Inches(0.5),
              "Every phase covered, plus end-to-end.",
              font_size=18, color=ACCENT_2, align=PP_ALIGN.CENTER)

    _footer(slide, page, total)


def make_techstack_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "10  ·  Toolbox", "Tech Stack")

    sections = [
        ("Backend",  ACCENT,   ["Python", "FastAPI", "Pydantic"]),
        ("NLP",      ACCENT_2, ["spaCy", "Transformers", "Tesseract"]),
        ("Storage",  WARN,     ["PostgreSQL", "MinIO"]),
        ("DevOps",   ACCENT,   ["Docker", "pytest"]),
    ]

    y = Inches(2.4)
    for title, color, items in sections:
        _add_text(slide, Inches(0.7), y, Inches(2.5), Inches(0.5),
                  title, font_size=20, bold=True, color=color)
        x = Inches(3.4)
        for item in items:
            c = _add_chip(slide, x, y + Inches(0.05), item, color=color)
            x += c.width + Inches(0.2)
        y += Inches(0.95)

    _footer(slide, page, total)


def make_challenges_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "11  ·  Reflection", "Challenges")

    _add_bullets(slide, Inches(0.9), Inches(2.4), Inches(12), Inches(5), [
        "Messy real-world layouts",
        "OCR noise",
        "NER models disagreeing",
        "Knowing when not to trust the model",
    ], font_size=26, line_spacing=1.7)

    _footer(slide, page, total)


def make_future_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _slide_header(slide, "12  ·  Roadmap", "Future Work")

    _add_bullets(slide, Inches(0.9), Inches(2.4), Inches(12), Inches(5), [
        "More document types",
        "Multilingual support",
        "Active learning from reviewer corrections",
        "Async batch processing",
    ], font_size=26, line_spacing=1.7)

    _footer(slide, page, total)


def make_thanks_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(0), Inches(0), Inches(0.3),
                                 Inches(7.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    _add_text(slide, Inches(0.9), Inches(2.8), Inches(11), Inches(1.6),
              "Thank you.", font_size=96, bold=True, color=TEXT)

    _add_text(slide, Inches(0.9), Inches(4.5), Inches(11), Inches(0.7),
              "Questions?",
              font_size=28, color=ACCENT_2)


# ---------------------------------------------------------------------------
# Driver
# ---------------------------------------------------------------------------

def build():
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)

    builders = [
        make_title_slide,
        make_problem_slide,
        make_features_slide,
        make_architecture_slide,
        make_pipeline_slide,
        make_phase_details_slide,
        make_schemas_slide,
        make_lifecycle_slide,
        make_results_slide,
        make_testing_slide,
        make_techstack_slide,
        make_challenges_slide,
        make_future_slide,
        make_thanks_slide,
    ]
    total = len(builders)

    builders[0](prs, total)
    for i, fn in enumerate(builders[1:], start=2):
        fn(prs, i, total)

    prs.save(OUTPUT)
    print(f"Saved -> {OUTPUT}")


if __name__ == "__main__":
    build()
