"""Phase 1 — Document ingestion.

Accepts PDF, DOCX, TXT, and image files. Extracts plain text. Falls back to
OCR (Tesseract) for image-only PDFs and image uploads.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from typing import List

import pdfplumber
from docx import Document as DocxDocument
from PIL import Image
import pytesseract

logger = logging.getLogger("docxtract.ingestion")

# Heuristic: if a parsed PDF page yields fewer chars than this, treat it as
# image-only and OCR it instead.
_OCR_FALLBACK_CHAR_THRESHOLD = 30

IMAGE_EXTENSIONS = frozenset({"png", "jpg", "jpeg", "tiff", "bmp", "webp"})
TEXT_EXTENSIONS = frozenset({"pdf", "docx", "doc", "txt", *IMAGE_EXTENSIONS})


@dataclass
class IngestionResult:
    """The output of phase 1: raw text plus per-source metadata."""

    text: str
    pages: List[str] = field(default_factory=list)
    used_ocr: bool = False
    source_format: str = ""
    char_count: int = 0


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def ingest(file_bytes: bytes, filename: str) -> IngestionResult:
    """Dispatch to the right parser based on file extension."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "txt":
        return _ingest_txt(file_bytes)
    if ext in {"docx", "doc"}:
        return _ingest_docx(file_bytes)
    if ext == "pdf":
        return _ingest_pdf(file_bytes)
    if ext in IMAGE_EXTENSIONS:
        return _ingest_image(file_bytes, ext)

    raise ValueError(f"Unsupported extension '.{ext}' for ingestion.")


# ---------------------------------------------------------------------------
# Per-format parsers
# ---------------------------------------------------------------------------

def _ingest_txt(file_bytes: bytes) -> IngestionResult:
    # Try utf-8 first, then latin-1 as a wide-net fallback.
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")
    return IngestionResult(
        text=text,
        pages=[text],
        source_format="txt",
        char_count=len(text),
    )


def _ingest_docx(file_bytes: bytes) -> IngestionResult:
    doc = DocxDocument(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Tables — extract cells row by row
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    text = "\n".join(paragraphs)
    return IngestionResult(
        text=text,
        pages=[text],
        source_format="docx",
        char_count=len(text),
    )


def _ingest_pdf(file_bytes: bytes) -> IngestionResult:
    """Parse PDF. If a page is image-only (very few chars), OCR it."""
    pages: List[str] = []
    used_ocr = False

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            page_text = (page.extract_text() or "").strip()

            if len(page_text) < _OCR_FALLBACK_CHAR_THRESHOLD:
                # Render page to image and OCR it.
                try:
                    page_image = page.to_image(resolution=200).original
                    ocr_text = pytesseract.image_to_string(page_image).strip()
                    if ocr_text:
                        page_text = ocr_text
                        used_ocr = True
                        logger.info("OCR fallback used on PDF page %d", page_num)
                except Exception:
                    logger.exception("OCR failed on PDF page %d", page_num)

            pages.append(page_text)

    text = "\n\n".join(pages)
    return IngestionResult(
        text=text,
        pages=pages,
        used_ocr=used_ocr,
        source_format="pdf",
        char_count=len(text),
    )


def _ingest_image(file_bytes: bytes, ext: str) -> IngestionResult:
    """OCR an uploaded image directly."""
    try:
        img = Image.open(io.BytesIO(file_bytes))
        # Convert to RGB to be safe for any mode
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        text = pytesseract.image_to_string(img).strip()
    except Exception as exc:
        logger.exception("Image ingestion failed")
        raise ValueError(f"Could not OCR image: {exc}") from exc

    return IngestionResult(
        text=text,
        pages=[text],
        used_ocr=True,
        source_format=ext,
        char_count=len(text),
    )
